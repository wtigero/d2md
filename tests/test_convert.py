import importlib
import sys
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pytest

from d2md.convert import (
    DEVICE_CHOICES,
    ConversionError,
    _docling_device,
    _ensure_device_available,
    _normalize_device,
    _configure_docling_for_platform,
    backend_for,
    convert,
)
from d2md.page_markers import format_pdf_content_span, pdf_content


convert_module = importlib.import_module("d2md.convert")


def _install_fake_docling(monkeypatch):
    captured = {"factory_calls": []}

    class PdfPipelineOptions:
        def __init__(self):
            self.do_ocr = None
            self.accelerator_options = None

    class FormatOption:
        def __init__(self, pipeline_options):
            self.pipeline_options = pipeline_options

    class DocumentConverter:
        def __init__(self, format_options):
            captured["format_options"] = format_options

    class Factory:
        def create_options(self, *, kind, lang):
            options = SimpleNamespace(mode=None)
            captured["factory_calls"].append((kind, lang, options))
            return options

    def get_ocr_factory(*, allow_external_plugins):
        captured["allow_external_plugins"] = allow_external_plugins
        return Factory()

    class OcrMode:
        FULL_PAGE = object()

    modules = {
        "docling": ModuleType("docling"),
        "docling.datamodel": ModuleType("docling.datamodel"),
        "docling.models": ModuleType("docling.models"),
        "docling.datamodel.accelerator_options": ModuleType(
            "docling.datamodel.accelerator_options"
        ),
        "docling.datamodel.base_models": ModuleType(
            "docling.datamodel.base_models"
        ),
        "docling.datamodel.pipeline_options": ModuleType(
            "docling.datamodel.pipeline_options"
        ),
        "docling.document_converter": ModuleType(
            "docling.document_converter"
        ),
        "docling.models.factories": ModuleType("docling.models.factories"),
    }
    for name in ("docling", "docling.datamodel", "docling.models"):
        modules[name].__path__ = []

    accelerator = modules["docling.datamodel.accelerator_options"]
    accelerator.AcceleratorDevice = SimpleNamespace(
        AUTO="auto", CPU="cpu", CUDA="cuda", MPS="mps", XPU="xpu"
    )
    accelerator.AcceleratorOptions = lambda **values: SimpleNamespace(**values)
    modules["docling.datamodel.base_models"].InputFormat = SimpleNamespace(
        PDF="pdf", IMAGE="image"
    )
    pipeline = modules["docling.datamodel.pipeline_options"]
    pipeline.PdfPipelineOptions = PdfPipelineOptions
    pipeline.OcrMode = OcrMode
    converter = modules["docling.document_converter"]
    converter.DocumentConverter = DocumentConverter
    converter.ImageFormatOption = FormatOption
    converter.PdfFormatOption = FormatOption
    modules["docling.models.factories"].get_ocr_factory = get_ocr_factory

    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)
    monkeypatch.setattr(convert_module, "_ensure_device_available", lambda _: None)
    monkeypatch.setattr(
        convert_module, "_configure_docling_for_platform", lambda *args, **kwargs: None
    )
    monkeypatch.setattr("d2md.ocr.engine_for", lambda script: "vision")
    captured["ocr_mode"] = OcrMode.FULL_PAGE
    return captured


def test_docling_builder_disables_ocr_without_explicit_ocr(monkeypatch):
    captured = _install_fake_docling(monkeypatch)

    convert_module._build_docling_converter(
        "thai", False, "cpu", ocr_enabled=False
    )

    options = captured["format_options"]["pdf"].pipeline_options
    assert options.do_ocr is False
    assert captured["factory_calls"] == []


def test_docling_builder_enables_bundled_ocr_and_full_page_mode(monkeypatch):
    captured = _install_fake_docling(monkeypatch)

    convert_module._build_docling_converter(
        "thai", True, "cpu", ocr_enabled=True
    )

    options = captured["format_options"]["pdf"].pipeline_options
    assert options.do_ocr is True
    assert captured["allow_external_plugins"] is False
    kind, languages, ocr_options = captured["factory_calls"][0]
    assert kind == "ocrmac"
    assert languages == ["th-TH"]
    assert ocr_options.mode is captured["ocr_mode"]


def test_device_choices_are_stable():
    assert DEVICE_CHOICES == ("auto", "cpu", "cuda", "mps", "xpu")


@pytest.mark.parametrize("device", ("auto", "cpu", "cuda", "mps", "xpu"))
def test_normalize_device_accepts_public_values(device):
    assert _normalize_device(device) == device


def test_normalize_device_rejects_generic_gpu():
    with pytest.raises(ConversionError, match="unknown device 'gpu'"):
        _normalize_device("gpu")


@pytest.mark.parametrize(
    "device,torch_module,message",
    [
        (
            "cuda",
            SimpleNamespace(
                version=SimpleNamespace(cuda=None),
                cuda=SimpleNamespace(is_available=lambda: False),
            ),
            "CUDA was requested",
        ),
        (
            "mps",
            SimpleNamespace(
                backends=SimpleNamespace(
                    mps=SimpleNamespace(is_available=lambda: False)
                )
            ),
            "MPS was requested",
        ),
        (
            "xpu",
            SimpleNamespace(xpu=SimpleNamespace(is_available=lambda: False)),
            "XPU was requested",
        ),
    ],
)
def test_unavailable_explicit_accelerator_fails(device, torch_module, message):
    with pytest.raises(ConversionError, match=message):
        _ensure_device_available(device, torch_module=torch_module)


def test_cpu_and_auto_need_no_accelerator_runtime():
    _ensure_device_available("cpu", torch_module=object())
    _ensure_device_available("auto", torch_module=object())


def test_each_public_device_maps_to_matching_docling_enum():
    sentinel = object()
    accelerator_device = SimpleNamespace(
        AUTO=sentinel,
        CPU=object(),
        CUDA=object(),
        MPS=object(),
        XPU=object(),
    )

    for device in DEVICE_CHOICES:
        assert _docling_device(device, accelerator_device) is getattr(
            accelerator_device, device.upper()
        )


def test_via_docling_translates_late_device_rejection(monkeypatch, tmp_path):
    """An explicit device rejection must survive unknown-format routing."""

    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    docling = ModuleType("docling")
    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = AcceleratorDeviceNotAvailableError
    docling.exceptions = exceptions
    monkeypatch.setitem(sys.modules, "docling", docling)
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)

    unavailable = AcceleratorDeviceNotAvailableError("CUDA is unavailable")
    monkeypatch.setattr(
        convert_module,
        "_docling",
        lambda *args, **kwargs: (_ for _ in ()).throw(unavailable),
    )

    with pytest.raises(ConversionError, match="CUDA is unavailable") as raised:
        convert_module._via_docling(tmp_path / "document.unknown", device="cuda")

    assert raised.value.__cause__ is unavailable


def _docling_pdf_document(pages, content, nested_page=None):
    pytest.importorskip("docling_core")
    from docling_core.types.doc.base import BoundingBox, CoordOrigin, Size
    from docling_core.types.doc.document import (
        DoclingDocument,
        PageItem,
        ProvenanceItem,
    )
    from docling_core.types.doc.labels import DocItemLabel, GroupLabel

    bbox = BoundingBox(l=0, t=1, r=1, b=0, coord_origin=CoordOrigin.BOTTOMLEFT)
    document = DoclingDocument(name="pages")
    for page_number in pages:
        document.pages[page_number] = PageItem(
            page_no=page_number,
            size=Size(width=1, height=1),
        )
    group = document.add_group(label=GroupLabel.SECTION) if nested_page else None
    for page_number, text in content.items():
        document.add_text(
            label=DocItemLabel.TEXT,
            text=text,
            parent=group if page_number == nested_page else None,
            prov=ProvenanceItem(
                page_no=page_number,
                bbox=bbox,
                charspan=(0, len(text)),
            ),
        )
    return document


def _forbid_docling_export(monkeypatch, document):
    seen = []

    def unexpected_export(*args, **kwargs):
        seen.append((args, kwargs))
        raise AssertionError("PDF conversion must not export once per page")

    monkeypatch.setattr(
        type(document), "export_to_markdown", unexpected_export
    )
    return seen


@pytest.mark.parametrize(
    ("pages", "content", "expected"),
    [
        pytest.param(
            (1, 2),
            {
                1: "First page content",
                2: "    indented second-page content",
            },
            "<!-- Page number: 1 -->\n\nFirst page content\n\n"
            "<!-- Page number: 2 -->\n\n    indented second-page content\n",
            id="consecutive-populated-pages",
        ),
        pytest.param(
            (1, 2, 3),
            {1: "First page content", 3: "Third page content"},
            "<!-- Page number: 1 -->\n\nFirst page content\n\n"
            "<!-- Page number: 2 -->\n\n"
            "<!-- Page number: 3 -->\n\nThird page content\n",
            id="blank-middle-page",
        ),
        pytest.param(
            (1, 2),
            {2: "Second page content"},
            "<!-- Page number: 1 -->\n\n"
            "<!-- Page number: 2 -->\n\nSecond page content\n",
            id="leading-blank-page",
        ),
    ],
)
def test_via_docling_marks_pdf_pages_in_source_order(
    monkeypatch, tmp_path, pages, content, expected
):
    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document(pages, content)
    seen = _forbid_docling_export(monkeypatch, document)

    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown == expected
    assert seen == []


def test_via_docling_keeps_nested_group_on_its_provenance_page(
    monkeypatch, tmp_path
):
    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document(
        (1, 2, 3),
        {1: "First page content", 3: "Nested third-page content"},
        nested_page=3,
    )
    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown == (
        "<!-- Page number: 1 -->\n\nFirst page content\n\n"
        "<!-- Page number: 2 -->\n\n"
        "<!-- Page number: 3 -->\n\nNested third-page content\n"
    )


def test_via_docling_marks_unsplittable_multi_page_text(monkeypatch, tmp_path):
    pytest.importorskip("docling_core")
    from docling_core.types.doc.base import BoundingBox, CoordOrigin
    from docling_core.types.doc.document import ProvenanceItem
    from docling_core.types.doc.labels import DocItemLabel

    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document((1, 2), {})
    bbox = BoundingBox(l=0, t=1, r=1, b=0, coord_origin=CoordOrigin.BOTTOMLEFT)
    text = "First-page text Second-page text"
    item = document.add_text(
        label=DocItemLabel.TEXT,
        text=text,
        prov=ProvenanceItem(page_no=1, bbox=bbox, charspan=(0, 16)),
    )
    item.prov.append(
        ProvenanceItem(page_no=2, bbox=bbox, charspan=(16, len(text)))
    )
    seen = _forbid_docling_export(monkeypatch, document)
    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown.count(text) == 1
    assert "<!-- Content spans pages: 1-2 -->" in markdown
    assert "<!-- Page number: 1 -->" in markdown
    assert "<!-- Page number: 2 -->" in markdown
    assert (
        markdown.index("<!-- Page number: 1 -->")
        < markdown.index("<!-- Content spans pages: 1-2 -->")
        < markdown.index(text)
        < markdown.index("<!-- Page number: 2 -->")
    )
    assert seen == []


def test_via_docling_marks_unsplittable_multi_page_table(monkeypatch, tmp_path):
    pytest.importorskip("docling_core")
    from docling_core.types.doc.base import BoundingBox, CoordOrigin
    from docling_core.types.doc.document import (
        ProvenanceItem,
        TableCell,
        TableData,
    )

    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document((1, 2), {})
    bbox = BoundingBox(l=0, t=1, r=1, b=0, coord_origin=CoordOrigin.BOTTOMLEFT)
    data = TableData(
        num_rows=2,
        num_cols=1,
        table_cells=[
            TableCell(
                text="Header",
                start_row_offset_idx=0,
                end_row_offset_idx=1,
                start_col_offset_idx=0,
                end_col_offset_idx=1,
                bbox=bbox,
            ),
            TableCell(
                text="Value",
                start_row_offset_idx=1,
                end_row_offset_idx=2,
                start_col_offset_idx=0,
                end_col_offset_idx=1,
                bbox=bbox,
            ),
        ],
    )
    table = document.add_table(
        data=data,
        prov=ProvenanceItem(page_no=1, bbox=bbox, charspan=(0, 0)),
    )
    table.prov.append(
        ProvenanceItem(page_no=2, bbox=bbox, charspan=(0, 0))
    )
    seen = _forbid_docling_export(monkeypatch, document)
    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown.count("Header") == 1
    assert markdown.count("Value") == 1
    assert "<!-- Content spans pages: 1-2 -->" in markdown
    assert "<!-- Page number: 1 -->" in markdown
    assert "<!-- Page number: 2 -->" in markdown
    assert (
        markdown.index("<!-- Page number: 1 -->")
        < markdown.index("<!-- Content spans pages: 1-2 -->")
        < markdown.index("Header")
        < markdown.index("<!-- Page number: 2 -->")
    )
    assert not hasattr(data.table_cells[0], "page_no")
    assert not hasattr(data.table_cells[0], "prov")
    assert seen == []


def test_via_docling_marks_atomic_segment_before_later_page_break(
    monkeypatch, tmp_path
):
    pytest.importorskip("docling_core")
    from docling_core.types.doc.base import BoundingBox, CoordOrigin
    from docling_core.types.doc.document import ProvenanceItem
    from docling_core.types.doc.labels import DocItemLabel, GroupLabel

    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document((1, 2, 3), {})
    bbox = BoundingBox(l=0, t=1, r=1, b=0, coord_origin=CoordOrigin.BOTTOMLEFT)
    group = document.add_group(label=GroupLabel.SECTION)
    text = "Atomic page one and two"
    item = document.add_text(
        label=DocItemLabel.TEXT,
        text=text,
        parent=group,
        prov=ProvenanceItem(page_no=1, bbox=bbox, charspan=(0, 11)),
    )
    item.prov.append(
        ProvenanceItem(page_no=2, bbox=bbox, charspan=(11, len(text)))
    )
    document.add_text(
        label=DocItemLabel.TEXT,
        text="Page three content",
        parent=group,
        prov=ProvenanceItem(page_no=3, bbox=bbox, charspan=(0, 18)),
    )
    seen = _forbid_docling_export(monkeypatch, document)
    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown == (
        "<!-- Page number: 1 -->\n\n"
        "<!-- Content spans pages: 1-2 -->\n\n"
        "Atomic page one and two\n\n"
        "<!-- Page number: 2 -->\n\n"
        "<!-- Page number: 3 -->\n\nPage three content\n"
    )
    assert markdown.count(text) == 1
    assert seen == []


def test_via_docling_falls_back_for_nested_group_trailing_break(
    monkeypatch, tmp_path
):
    pytest.importorskip("docling_core")
    from docling_core.types.doc.base import BoundingBox, CoordOrigin
    from docling_core.types.doc.document import ProvenanceItem
    from docling_core.types.doc.labels import DocItemLabel, GroupLabel

    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)
    document = _docling_pdf_document((1, 2, 3), {})
    bbox = BoundingBox(l=0, t=1, r=1, b=0, coord_origin=CoordOrigin.BOTTOMLEFT)
    outer = document.add_group(label=GroupLabel.SECTION)
    document.add_text(
        label=DocItemLabel.TEXT,
        text="Page one content",
        parent=outer,
        prov=ProvenanceItem(page_no=1, bbox=bbox, charspan=(0, 16)),
    )
    nested = document.add_group(label=GroupLabel.SECTION, parent=outer)
    document.add_text(
        label=DocItemLabel.TEXT,
        text="Nested page three content",
        parent=nested,
        prov=ProvenanceItem(page_no=3, bbox=bbox, charspan=(0, 25)),
    )
    seen = _forbid_docling_export(monkeypatch, document)
    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=document)
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    markdown = convert_module._via_docling(tmp_path / "report.pdf")

    assert markdown.count("Page one content") == 1
    assert markdown.count("Nested page three content") == 1
    assert "<!-- Content spans pages: 1,3 -->" in markdown
    assert convert_module.DOCLING_PAGE_BREAK not in markdown
    assert "#_#_DOCLING_DOC_PAGE_BREAK_" not in markdown
    assert markdown.count("<!-- Page number:") == 3
    assert (
        markdown.index("<!-- Page number: 1 -->")
        < markdown.index("<!-- Content spans pages: 1,3 -->")
        < markdown.index("Page one content")
        < markdown.index("Nested page three content")
        < markdown.index("<!-- Page number: 2 -->")
        < markdown.index("<!-- Page number: 3 -->")
    )
    assert seen == []


def test_pdf_content_ignores_page_metadata():
    markdown = (
        "<!-- Page number: 1 -->\n\n"
        "<!-- Content spans pages: 1-2 -->\n"
    )

    assert not pdf_content(markdown).strip()
    assert format_pdf_content_span([1, 3]) == "<!-- Content spans pages: 1,3 -->"


def test_via_docling_keeps_image_markdown_unchanged(monkeypatch, tmp_path):
    class AcceleratorDeviceNotAvailableError(RuntimeError):
        pass

    exceptions = ModuleType("docling.exceptions")
    exceptions.AcceleratorDeviceNotAvailableError = (
        AcceleratorDeviceNotAvailableError
    )
    monkeypatch.setitem(sys.modules, "docling.exceptions", exceptions)

    class Document:
        def export_to_markdown(self):
            return "Image description"

    converter = SimpleNamespace(
        convert=lambda path: SimpleNamespace(document=Document())
    )
    monkeypatch.setattr(
        convert_module, "_docling", lambda *args, **kwargs: converter
    )

    assert (
        convert_module._via_docling(tmp_path / "scan.png")
        == "Image description"
    )


def test_via_docling_reraises_unrelated_runtime_error(monkeypatch, tmp_path):
    # _via_docling imports docling before it reaches the patched _docling, so
    # this reaches the reraise path only where the optional profile is present.
    pytest.importorskip("docling")

    unrelated = RuntimeError("layout model failed")
    monkeypatch.setattr(
        convert_module,
        "_docling",
        lambda *args, **kwargs: (_ for _ in ()).throw(unrelated),
    )

    with pytest.raises(RuntimeError) as raised:
        convert_module._via_docling(tmp_path / "document.unknown")

    assert raised.value is unrelated


@pytest.mark.parametrize(
    "name,docling,expected",
    [
        ("report.pdf", False, "pypdfium2"),
        ("scan.png", False, "ocr"),
        ("report.pdf", True, "docling"),
        ("scan.png", True, "docling"),
        ("sheet.xlsx", False, "markitdown"),
        ("notes.txt", False, "plain"),
        ("mystery.bin", False, "markitdown"),
        ("mystery.bin", True, "markitdown"),
    ],
)
def test_explicit_routing(name, docling, expected):
    assert backend_for(Path(name), docling=docling) == expected


def test_windows_disables_optional_docling_model_compilation():
    engine = SimpleNamespace(compile_model=True)
    options = SimpleNamespace(
        layout_options=SimpleNamespace(engine_options=engine)
    )

    _configure_docling_for_platform(options, platform="nt")

    assert engine.compile_model is False


@pytest.mark.parametrize(
    "capability,expected",
    [
        ((6, 1), False),
        ((7, 0), True),
    ],
)
def test_docling_model_compilation_respects_cuda_capability(
    monkeypatch, capability, expected
):
    engine = SimpleNamespace(compile_model=True)
    options = SimpleNamespace(
        layout_options=SimpleNamespace(engine_options=engine)
    )
    cuda = SimpleNamespace(
        is_available=lambda: True,
        get_device_capability=lambda: capability,
    )
    monkeypatch.setitem(
        sys.modules,
        "torch",
        SimpleNamespace(cuda=cuda, version=SimpleNamespace(cuda="11.8")),
    )

    _configure_docling_for_platform(options, platform="posix")

    assert engine.compile_model is expected


def test_forced_cpu_does_not_apply_installed_pascal_limit(monkeypatch):
    engine = SimpleNamespace(compile_model=True)
    options = SimpleNamespace(
        layout_options=SimpleNamespace(engine_options=engine)
    )
    cuda = SimpleNamespace(
        is_available=lambda: True,
        get_device_capability=lambda: (6, 1),
    )
    monkeypatch.setitem(
        sys.modules,
        "torch",
        SimpleNamespace(cuda=cuda, version=SimpleNamespace(cuda="11.8")),
    )

    _configure_docling_for_platform(options, device="cpu", platform="posix")

    assert engine.compile_model is True


def test_pascal_cuda_disables_compilation(monkeypatch):
    engine = SimpleNamespace(compile_model=True)
    options = SimpleNamespace(
        layout_options=SimpleNamespace(engine_options=engine)
    )
    cuda = SimpleNamespace(
        is_available=lambda: True,
        get_device_capability=lambda: (6, 1),
    )
    monkeypatch.setitem(
        sys.modules,
        "torch",
        SimpleNamespace(cuda=cuda, version=SimpleNamespace(cuda="11.8")),
    )

    _configure_docling_for_platform(options, device="cuda", platform="posix")

    assert engine.compile_model is False


def test_converter_cache_is_isolated_by_every_device(monkeypatch):
    import importlib

    convert_module = importlib.import_module("d2md.convert")

    built = []

    def fake_build(script, force_ocr, device, ocr_enabled):
        converter = object()
        built.append((script, force_ocr, device, ocr_enabled, converter))
        return converter

    monkeypatch.setattr(convert_module, "_build_docling_converter", fake_build)
    convert_module._converters.clear()

    converters = {
        device: convert_module._docling("latin", device=device)
        for device in DEVICE_CHOICES
    }

    assert len({id(converter) for converter in converters.values()}) == 5
    for device, converter in converters.items():
        assert convert_module._docling("latin", device=device) is converter
    assert [
        (script, force, device, ocr_enabled)
        for script, force, device, ocr_enabled, _ in built
    ] == [
        ("latin", False, "auto", False),
        ("latin", False, "cpu", False),
        ("latin", False, "cuda", False),
        ("latin", False, "mps", False),
        ("latin", False, "xpu", False),
    ]


def test_docling_cache_separates_ocr_state(monkeypatch):
    built = []

    def fake_build(script, force_ocr, device, ocr_enabled):
        converter = object()
        built.append((script, force_ocr, device, ocr_enabled, converter))
        return converter

    monkeypatch.setattr(convert_module, "_build_docling_converter", fake_build)
    convert_module._converters.clear()

    without = convert_module._docling("thai", device="cpu", ocr_enabled=False)
    with_ocr = convert_module._docling("thai", device="cpu", ocr_enabled=True)

    assert without is not with_ocr
    assert [
        (script, force, device, ocr_enabled)
        for script, force, device, ocr_enabled, _ in built
    ] == [
        ("thai", False, "cpu", False),
        ("thai", False, "cpu", True),
    ]


def test_empty_output_is_an_error_not_a_success(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_text("")
    with pytest.raises(ConversionError):
        convert(empty)


def test_short_output_is_an_error(tmp_path):
    tiny = tmp_path / "tiny.txt"
    tiny.write_text("hello")
    with pytest.raises(ConversionError):
        convert(tiny)


def test_plain_text_passes_through(tmp_path):
    src = tmp_path / "note.txt"
    body = "ทดสอบภาษาไทย น้ำจำกัด ระบบอ่านไฟล์ทดสอบ"
    src.write_text(body, encoding="utf-8")
    result = convert(src)
    assert result.backend == "plain"
    assert result.markdown == body


@pytest.mark.parametrize("suffix", (".xlsx", ".xls"))
def test_spreadsheet_empty_cells_render_as_blank_table_cells(tmp_path, suffix):
    source = tmp_path / f"fields{suffix}"
    if suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Field", "Value"])
        sheet.append(["Missing", None])
        sheet.append(["Literal", "NaN"])
        sheet.append(["Text", "Contains NaN text"])
        sheet.append(["Formula error", "#REF!"])
        workbook.save(source)
    else:
        import xlwt

        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Sheet1")
        sheet.write(0, 0, "Field")
        sheet.write(0, 1, "Value")
        sheet.write(1, 0, "Missing")
        sheet.write(2, 0, "Literal")
        sheet.write(2, 1, "NaN")
        sheet.write(3, 0, "Text")
        sheet.write(3, 1, "Contains NaN text")
        sheet.write(4, 0, "Formula error")
        sheet.write(4, 1, xlwt.Formula("1/0"))
        workbook.save(str(source))

    result = convert(source)
    markdown = result.markdown

    assert result.backend == "markitdown"
    missing_row = next(
        line for line in markdown.splitlines() if line.startswith("| Missing |")
    )
    assert missing_row == "| Missing |  |"
    formula_row = next(
        line for line in markdown.splitlines() if line.startswith("| Formula error |")
    )
    assert formula_row == "| Formula error |  |"
    assert "| Literal | NaN |" in markdown
    assert "Contains NaN text" in markdown


def test_legacy_thai_text_file_survives(tmp_path):
    src = tmp_path / "legacy.txt"
    body = "ทดสอบภาษาไทย น้ำจำกัด ระบบอ่านไฟล์ทดสอบ"
    src.write_bytes(body.encode("cp874"))
    assert convert(src).markdown == body
